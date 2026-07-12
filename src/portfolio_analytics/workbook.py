from __future__ import annotations

from io import BytesIO
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import PortfolioWorkbook, ValidationMessage


SHEET_PORTFOLIO_COMPOSITION = "Portfolio Composition"
SHEET_HISTORICAL_PRICES = "Historical Prices"
SHEET_TRANSACTIONS = "Transactions"
REQUIRED_SHEETS = [
    SHEET_PORTFOLIO_COMPOSITION,
    SHEET_HISTORICAL_PRICES,
    SHEET_TRANSACTIONS,
]

TRANSACTION_EXCEL_HEADERS = [
    "Broker",
    "Product",
    "ISIN",
    "Date",
    "Exchange",
    "Action",
    "Quantity",
    "Price in \u20ac",
    "Cash Flow in \u20ac",
    "Fee in \u20ac",
]
TRANSACTION_COLUMNS = [
    "broker",
    "product",
    "isin",
    "date",
    "exchange",
    "action",
    "quantity",
    "price_eur",
    "cash_flow_eur",
    "fee_eur",
]
PRICE_HEADERS = ["Product", "ISIN", "Date", "Price"]
RISK_FREE_GROUP_HEADER = "Risk-Free Government Bonds"
RISK_FREE_EXCEL_HEADERS = [
    "Year",
    "Country",
    "1Y-Yield",
    "3Y-Yield",
    "5Y-Yield",
    "10Y-Yield",
    "20Y-Yield",
    "30Y-Yield",
]
RISK_FREE_COLUMNS = [
    "year",
    "country",
    "yield_1y",
    "yield_3y",
    "yield_5y",
    "yield_10y",
    "yield_20y",
    "yield_30y",
]
COMPOSITION_TARGET_HEADER = "Target Weight in Portfolio"
COMPOSITION_TOLERANCE_HEADER = "Weight Deviation Tolerance"
SUPPORTED_ACTIONS = {"Buy", "Sell", "Transfer"}


def _clean_text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _canonical(value) -> str:
    return " ".join(_clean_text(value).lower().replace("\u00a0", " ").split())


def _normalize_isin(value) -> str:
    text = _clean_text(value).upper().replace(" ", "")
    if text in {"", "NAN", "NONE"}:
        return ""
    return text


def make_asset_id(product, isin) -> str:
    isin_text = _normalize_isin(isin)
    product_text = _clean_text(product)
    return isin_text or product_text


def _message(messages: list[ValidationMessage], severity: str, sheet: str, message: str, rows: Iterable[int] | None = None) -> None:
    row_text = ""
    if rows:
        row_text = ", ".join(str(r) for r in sorted(set(rows))[:20])
    messages.append(ValidationMessage(severity=severity, sheet=sheet, message=message, rows=row_text))


def _parse_number(value) -> float:
    if value is None or pd.isna(value):
        return np.nan
    if isinstance(value, str):
        text = value.strip().replace("\u20ac", "").replace(",", "")
        if text == "":
            return np.nan
        return pd.to_numeric(text, errors="coerce")
    return pd.to_numeric(value, errors="coerce")


def _parse_percent(value) -> tuple[float, bool]:
    if value is None or pd.isna(value):
        return np.nan, False
    if isinstance(value, str):
        text = value.strip().replace(",", ".")
        if not text:
            return np.nan, False
        if text.endswith("%"):
            parsed = pd.to_numeric(text[:-1], errors="coerce")
            return (float(parsed) / 100.0 if pd.notna(parsed) else np.nan), False
        parsed = pd.to_numeric(text, errors="coerce")
    else:
        parsed = pd.to_numeric(value, errors="coerce")
    if pd.isna(parsed):
        return np.nan, False
    parsed = float(parsed)
    if 1.0 < abs(parsed) <= 100.0:
        return parsed / 100.0, True
    return parsed, False


def _normalize_action(value) -> str:
    text = _canonical(value)
    if text in {"buy", "bought", "purchase", "purchased"}:
        return "Buy"
    if text in {"sell", "sold", "sale"}:
        return "Sell"
    if text in {"transfer", "transfer in", "transfer out"}:
        return "Transfer"
    return _clean_text(value)


def _read_excel_raw(file_obj) -> tuple[pd.ExcelFile, dict[str, pd.DataFrame]]:
    xls = pd.ExcelFile(file_obj)
    raw = {sheet: xls.parse(sheet, header=None) for sheet in xls.sheet_names}
    return xls, raw


def parse_workbook(file_obj) -> PortfolioWorkbook:
    xls, raw_sheets = _read_excel_raw(file_obj)
    messages: list[ValidationMessage] = []

    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in raw_sheets]
    for sheet in missing:
        _message(messages, "Error", sheet, f"Required sheet '{sheet}' is missing.")
    extra = [sheet for sheet in xls.sheet_names if sheet not in REQUIRED_SHEETS]
    if extra:
        _message(messages, "Warning", "Workbook", "Input workbook has extra sheets; calculations use only the three required input sheets.")

    transactions = (
        _parse_transactions(raw_sheets[SHEET_TRANSACTIONS], messages)
        if SHEET_TRANSACTIONS in raw_sheets
        else _empty_transactions()
    )
    prices = (
        _parse_prices(raw_sheets[SHEET_HISTORICAL_PRICES], messages)
        if SHEET_HISTORICAL_PRICES in raw_sheets
        else _empty_prices()
    )
    risk_free_rates = (
        _parse_risk_free_rates(raw_sheets[SHEET_HISTORICAL_PRICES], messages)
        if SHEET_HISTORICAL_PRICES in raw_sheets
        else _empty_risk_free_rates()
    )
    composition = (
        _parse_composition(raw_sheets[SHEET_PORTFOLIO_COMPOSITION], messages)
        if SHEET_PORTFOLIO_COMPOSITION in raw_sheets
        else _empty_composition()
    )

    _cross_validate(transactions, prices, composition, messages)

    return PortfolioWorkbook(
        transactions=transactions,
        prices=prices,
        composition=composition,
        validation=messages,
        workbook_sheets=xls.sheet_names,
        risk_free_rates=risk_free_rates,
    )


def _empty_transactions() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "broker",
            "product",
            "isin",
            "asset_id",
            "date",
            "exchange",
            "action",
            "quantity",
            "price_eur",
            "cash_flow_eur",
            "fee_eur",
            "source_row",
        ]
    )


def _empty_prices() -> pd.DataFrame:
    return pd.DataFrame(columns=["product", "isin", "asset_id", "date", "price", "source_row", "block"])


def _empty_risk_free_rates() -> pd.DataFrame:
    return pd.DataFrame(columns=RISK_FREE_COLUMNS)


def _empty_composition() -> pd.DataFrame:
    return pd.DataFrame(columns=["year", "product", "asset_id", "target_weight", "tolerance", "source_row", "target_column"])


def _parse_transactions(raw: pd.DataFrame, messages: list[ValidationMessage]) -> pd.DataFrame:
    sheet = SHEET_TRANSACTIONS
    if raw.shape[0] < 2 or raw.shape[1] < 10:
        _message(messages, "Error", sheet, "Transactions sheet must contain A1:J2 with the All Transactions section and exact headers.")
        return _empty_transactions()

    group_header = _clean_text(raw.iat[0, 0])
    if group_header != "All Transactions":
        _message(messages, "Error", sheet, "Cell A1 must be 'All Transactions'.")

    row2 = [_clean_text(raw.iat[1, idx]) if idx < raw.shape[1] else "" for idx in range(10)]
    if row2 != TRANSACTION_EXCEL_HEADERS:
        _message(messages, "Error", sheet, "Row 2 headers must exactly match the required A:J Transactions labels.")

    body = raw.iloc[2:, :10].copy()
    body.columns = TRANSACTION_COLUMNS
    body["source_row"] = body.index + 1
    body = body.dropna(how="all").copy()
    if body.empty:
        _message(messages, "Warning", sheet, "Transactions sheet contains no data rows.")
        return _empty_transactions()

    for col in ["broker", "product", "isin", "exchange", "action"]:
        body[col] = body[col].map(_clean_text)
    body["isin"] = body["isin"].map(_normalize_isin)
    body["asset_id"] = [make_asset_id(p, i) for p, i in zip(body["product"], body["isin"])]
    body["date"] = pd.to_datetime(body["date"], errors="coerce").dt.normalize()
    body["action"] = body["action"].map(_normalize_action)
    for col in ["quantity", "price_eur", "cash_flow_eur", "fee_eur"]:
        body[col] = body[col].map(_parse_number).astype(float)
    body["fee_eur"] = body["fee_eur"].fillna(0.0)

    _flag_rows(body, body["date"].isna(), messages, "Error", sheet, "Transaction dates must parse as dates.")
    _flag_rows(body, body["asset_id"].eq(""), messages, "Error", sheet, "Each transaction row needs an ISIN or Product.")

    unsupported = ~body["action"].isin(SUPPORTED_ACTIONS)
    _flag_rows(body, unsupported, messages, "Warning", sheet, "Unsupported actions are preserved but ignored by calculations.")

    cash_actions = body["action"].isin(["Buy", "Sell"])
    _flag_rows(body, cash_actions & body["quantity"].isna(), messages, "Error", sheet, "Buy/Sell rows require numeric Quantity.")
    _flag_rows(body, cash_actions & body["price_eur"].isna(), messages, "Warning", sheet, "Buy/Sell rows should include numeric Price in EUR.")
    _flag_rows(body, cash_actions & body["cash_flow_eur"].isna(), messages, "Error", sheet, "Buy/Sell rows require numeric Cash Flow in EUR.")
    _flag_rows(body, cash_actions & (body["quantity"] < 0), messages, "Warning", sheet, "Quantity should be entered as a positive amount; action controls direction.")
    _flag_rows(body, cash_actions & (body["cash_flow_eur"] < 0), messages, "Warning", sheet, "Cash Flow in EUR should be entered as a positive amount.")

    duplicates = body.duplicated(
        subset=["broker", "product", "isin", "date", "action", "quantity", "price_eur", "cash_flow_eur", "fee_eur"],
        keep=False,
    )
    _flag_rows(body, duplicates, messages, "Warning", sheet, "Potential duplicate transaction rows detected.")

    valid = body.dropna(subset=["date"]).copy()
    valid["quantity"] = valid["quantity"].fillna(0.0)
    valid["price_eur"] = valid["price_eur"].fillna(0.0)
    valid["cash_flow_eur"] = valid["cash_flow_eur"].fillna(0.0)

    signed = np.select(
        [valid["action"].eq("Buy"), valid["action"].eq("Sell")],
        [valid["quantity"].abs(), -valid["quantity"].abs()],
        default=0.0,
    )
    valid["_signed_quantity"] = signed
    quantity_check = valid.sort_values(["asset_id", "date", "source_row"]).copy()
    quantity_check["cumulative_quantity"] = quantity_check.groupby("asset_id")["_signed_quantity"].cumsum()
    _flag_rows(
        quantity_check,
        quantity_check["cumulative_quantity"] < -1e-9,
        messages,
        "Warning",
        sheet,
        "Cumulative quantity becomes negative after a sell.",
    )

    return valid[_empty_transactions().columns].sort_values(["date", "source_row"]).reset_index(drop=True)


def _parse_prices(raw: pd.DataFrame, messages: list[ValidationMessage]) -> pd.DataFrame:
    sheet = SHEET_HISTORICAL_PRICES
    if raw.shape[0] < 2:
        _message(messages, "Error", sheet, "Historical Prices sheet must have row 1 group headers and row 2 block headers.")
        return _empty_prices()

    header_row = [_clean_text(v) for v in raw.iloc[1].tolist()]
    block_starts = [
        col
        for col in range(max(0, len(header_row) - len(PRICE_HEADERS) + 1))
        if header_row[col : col + len(PRICE_HEADERS)] == PRICE_HEADERS
    ]
    for col in block_starts:
        spacer_col = col + len(PRICE_HEADERS)
        if spacer_col < len(header_row) and _clean_text(raw.iat[1, spacer_col]) and header_row[spacer_col : spacer_col + len(RISK_FREE_EXCEL_HEADERS)] != RISK_FREE_EXCEL_HEADERS:
            _message(messages, "Warning", sheet, f"Expected a blank spacer column after price block starting at {get_column_letter(col + 1)}.")

    if not block_starts:
        _message(messages, "Error", sheet, "No valid Product/ISIN/Date/Price price blocks were found.")
        return _empty_prices()

    rows: list[dict] = []
    for block_number, start_col in enumerate(block_starts, start=1):
        group_header = _clean_text(raw.iat[0, start_col]) if raw.shape[0] else ""
        if not group_header:
            _message(messages, "Info", sheet, f"Price block {block_number} has no row-1 product/group header.")
        for row_idx in range(2, raw.shape[0]):
            product = _clean_text(raw.iat[row_idx, start_col]) if start_col < raw.shape[1] else ""
            isin = _normalize_isin(raw.iat[row_idx, start_col + 1]) if start_col + 1 < raw.shape[1] else ""
            date_raw = raw.iat[row_idx, start_col + 2] if start_col + 2 < raw.shape[1] else np.nan
            price_raw = raw.iat[row_idx, start_col + 3] if start_col + 3 < raw.shape[1] else np.nan
            if not product and not isin and pd.isna(date_raw) and pd.isna(price_raw):
                continue
            date = pd.to_datetime(date_raw, errors="coerce")
            price = _parse_number(price_raw)
            excel_row = row_idx + 1
            if not product and not isin:
                _message(messages, "Error", sheet, "Price rows need Product or ISIN.", [excel_row])
                continue
            if pd.isna(date):
                _message(messages, "Error", sheet, "Price dates must parse as dates.", [excel_row])
                continue
            if pd.isna(price):
                _message(messages, "Error", sheet, "Price rows need a numeric Price.", [excel_row])
                continue
            if float(price) <= 0:
                _message(messages, "Error", sheet, "Prices must be positive.", [excel_row])
                continue
            rows.append(
                {
                    "product": product or group_header,
                    "isin": isin,
                    "asset_id": make_asset_id(product or group_header, isin),
                    "date": pd.Timestamp(date).normalize(),
                    "price": float(price),
                    "source_row": excel_row,
                    "block": block_number,
                }
            )

    out = pd.DataFrame(rows, columns=_empty_prices().columns)
    if out.empty:
        _message(messages, "Error", sheet, "No usable historical price rows were found.")
        return out

    duplicate_mask = out.duplicated(subset=["asset_id", "date"], keep=False)
    _flag_rows(out, duplicate_mask, messages, "Warning", sheet, "Duplicate price rows for the same asset/date were found; the last row is used.")
    out = out.drop_duplicates(subset=["asset_id", "date"], keep="last")

    for asset_id, asset_prices in out.sort_values("date").groupby("asset_id"):
        gaps = asset_prices["date"].diff().dt.days.dropna()
        if not gaps.empty and gaps.max() > 14:
            _message(messages, "Info", sheet, f"{asset_id} has price-history gaps up to {int(gaps.max())} days.")

    return out.sort_values(["asset_id", "date"]).reset_index(drop=True)


def _parse_risk_free_rates(
    raw: pd.DataFrame, messages: list[ValidationMessage]
) -> pd.DataFrame:
    """Parse the optional government-bond yield block from Historical Prices."""
    sheet = SHEET_HISTORICAL_PRICES
    if raw.shape[0] < 2:
        _message(messages, "Warning", sheet, "No risk-free-rate block found. Risk-adjusted metrics use 0%.")
        return _empty_risk_free_rates()

    headers = [_clean_text(v) for v in raw.iloc[1].tolist()]
    starts = [
        col
        for col in range(max(0, len(headers) - len(RISK_FREE_EXCEL_HEADERS) + 1))
        if headers[col : col + len(RISK_FREE_EXCEL_HEADERS)] == RISK_FREE_EXCEL_HEADERS
    ]
    if not starts:
        _message(messages, "Warning", sheet, "No risk-free-rate block found. Risk-adjusted metrics use 0%.")
        return _empty_risk_free_rates()

    start = starts[0]
    group_header = _clean_text(raw.iat[0, start])
    if group_header != RISK_FREE_GROUP_HEADER:
        _message(
            messages,
            "Info",
            sheet,
            "Risk-free yields were detected from their header sequence (the group header is optional).",
        )
    if len(starts) > 1:
        _message(messages, "Warning", sheet, "Multiple risk-free-rate blocks were found; only the first block is used.")

    rows: list[dict] = []
    yield_columns = RISK_FREE_COLUMNS[2:]
    for row_idx in range(2, raw.shape[0]):
        values = [raw.iat[row_idx, start + offset] if start + offset < raw.shape[1] else np.nan for offset in range(8)]
        if all(pd.isna(value) or _clean_text(value) == "" for value in values):
            continue
        excel_row = row_idx + 1
        year_value = pd.to_numeric(values[0], errors="coerce")
        country = _clean_text(values[1])
        if pd.isna(year_value) or int(year_value) != float(year_value):
            _message(messages, "Error", sheet, "Risk-free-rate Year must be an integer.", [excel_row])
            continue
        if not country:
            _message(messages, "Warning", sheet, "Risk-free-rate rows should include a Country.", [excel_row])

        row = {"year": int(year_value), "country": country, "source_row": excel_row}
        valid_yields = 0
        for value, column, label in zip(values[2:], yield_columns, RISK_FREE_EXCEL_HEADERS[2:]):
            parsed, interpreted = _parse_percent(value)
            if pd.notna(parsed):
                valid_yields += 1
            elif not (pd.isna(value) or _clean_text(value) == ""):
                _message(messages, "Error", sheet, f"Risk-free {label} must be numeric.", [excel_row])
            if interpreted:
                _message(messages, "Info", sheet, f"Interpreted {label} in row {excel_row} as a percentage.")
            row[column] = parsed
        if valid_yields == 0:
            _message(messages, "Warning", sheet, "Risk-free-rate rows need at least one yield.", [excel_row])
            continue
        rows.append(row)

    result = pd.DataFrame(rows)
    if result.empty:
        _message(messages, "Warning", sheet, "No usable risk-free-rate rows were found. Risk-adjusted metrics use 0%.")
        return _empty_risk_free_rates()
    duplicates = result.duplicated(subset=["year", "country"], keep=False)
    _flag_rows(result, duplicates, messages, "Warning", sheet, "Duplicate risk-free rate rows were found; the last row is used.")
    result = result.drop_duplicates(subset=["year", "country"], keep="last")
    return result[RISK_FREE_COLUMNS].sort_values(["year", "country"]).reset_index(drop=True)


def _parse_composition(raw: pd.DataFrame, messages: list[ValidationMessage]) -> pd.DataFrame:
    sheet = SHEET_PORTFOLIO_COMPOSITION
    if raw.shape[0] < 2 or raw.shape[1] < 3:
        _message(messages, "Error", sheet, "Portfolio Composition sheet must contain row 1 title and row 2 dynamic headers.")
        return _empty_composition()

    title = _clean_text(raw.iat[0, 0])
    if title != SHEET_PORTFOLIO_COMPOSITION:
        _message(messages, "Warning", sheet, "Cell A1 should be 'Portfolio Composition'.")

    headers = [_clean_text(v) for v in raw.iloc[1].tolist()]
    if headers[0] != "Year":
        _message(messages, "Error", sheet, "Cell A2 must be 'Year'.")

    tolerance_candidates = [idx for idx, h in enumerate(headers) if h == COMPOSITION_TOLERANCE_HEADER]
    if not tolerance_candidates:
        _message(messages, "Error", sheet, "A final Weight Deviation Tolerance column is required.")
        return _empty_composition()
    tolerance_col = tolerance_candidates[-1]

    pairs: list[tuple[int, int]] = []
    col = 1
    while col + 1 < tolerance_col:
        instrument_header = headers[col]
        weight_header = headers[col + 1]
        if not instrument_header.startswith("Instrument") or weight_header != COMPOSITION_TARGET_HEADER:
            _message(messages, "Warning", sheet, f"Unexpected composition headers near {get_column_letter(col + 1)}.")
        pairs.append((col, col + 1))
        col += 2

    if not pairs:
        _message(messages, "Error", sheet, "No Instrument / Target Weight pairs were detected.")
        return _empty_composition()

    rows: list[dict] = []
    for row_idx in range(2, raw.shape[0]):
        excel_row = row_idx + 1
        row_values = raw.iloc[row_idx].tolist()
        if all(pd.isna(v) or _clean_text(v) == "" for v in row_values):
            continue
        year_value = pd.to_numeric(raw.iat[row_idx, 0], errors="coerce")
        if pd.isna(year_value):
            _message(messages, "Error", sheet, "Year must be an integer.", [excel_row])
            continue
        year = int(year_value)
        tolerance, interpreted_tolerance = _parse_percent(raw.iat[row_idx, tolerance_col])
        if pd.isna(tolerance):
            tolerance = 0.0
        if interpreted_tolerance:
            _message(messages, "Info", sheet, f"Interpreted tolerance in row {excel_row} as a percentage.")

        year_weights: list[float] = []
        for instrument_col, weight_col in pairs:
            product = _clean_text(raw.iat[row_idx, instrument_col]) if instrument_col < raw.shape[1] else ""
            weight_raw = raw.iat[row_idx, weight_col] if weight_col < raw.shape[1] else np.nan
            weight, interpreted_weight = _parse_percent(weight_raw)
            if not product and pd.isna(weight):
                continue
            if not product:
                _message(messages, "Error", sheet, "Target allocation rows need an Instrument name.", [excel_row])
                continue
            if pd.isna(weight):
                _message(messages, "Error", sheet, "Target weights must be numeric.", [excel_row])
                continue
            if interpreted_weight:
                _message(messages, "Info", sheet, f"Interpreted target weight in row {excel_row} as a percentage.")
            year_weights.append(float(weight))
            rows.append(
                {
                    "year": year,
                    "product": product,
                    "asset_id": product,
                    "target_weight": float(weight),
                    "tolerance": float(tolerance),
                    "source_row": excel_row,
                    "target_column": get_column_letter(weight_col + 1),
                }
            )
        if year_weights and abs(sum(year_weights) - 1.0) > 0.005:
            _message(messages, "Warning", sheet, f"Target weights for {year} sum to {sum(year_weights):.2%}, not 100%.", [excel_row])

    out = pd.DataFrame(rows, columns=_empty_composition().columns)
    if out.empty:
        _message(messages, "Warning", sheet, "Portfolio Composition contains no target allocation rows.")
    return out.sort_values(["year", "source_row", "target_column"]).reset_index(drop=True)


def _flag_rows(
    df: pd.DataFrame,
    mask: pd.Series,
    messages: list[ValidationMessage],
    severity: str,
    sheet: str,
    message: str,
) -> None:
    if mask.any():
        rows = df.loc[mask, "source_row"].dropna().astype(int).tolist() if "source_row" in df.columns else []
        _message(messages, severity, sheet, message, rows)


def _cross_validate(
    transactions: pd.DataFrame,
    prices: pd.DataFrame,
    composition: pd.DataFrame,
    messages: list[ValidationMessage],
) -> None:
    if transactions.empty:
        _message(messages, "Error", SHEET_TRANSACTIONS, "No usable transactions were found.")
    if prices.empty:
        _message(messages, "Error", SHEET_HISTORICAL_PRICES, "No usable historical prices were found.")

    tx_assets = set(transactions["asset_id"].dropna()) if not transactions.empty else set()
    price_assets = set(prices["asset_id"].dropna()) if not prices.empty else set()
    composition_assets = set(composition["asset_id"].dropna()) if not composition.empty else set()

    missing_prices = tx_assets - price_assets
    if missing_prices:
        _message(messages, "Warning", "Cross-sheet", "Assets with transactions but no historical prices: " + ", ".join(sorted(missing_prices)))
    price_without_tx = price_assets - tx_assets
    if price_without_tx:
        _message(messages, "Info", "Cross-sheet", "Assets with prices but no transactions: " + ", ".join(sorted(price_without_tx)))

    known_products = set()
    if not transactions.empty:
        known_products.update(transactions["product"].dropna().map(_clean_text))
    if not prices.empty:
        known_products.update(prices["product"].dropna().map(_clean_text))
    composition_products = set(composition["product"].dropna().map(_clean_text)) if not composition.empty else set()
    assigned_tx_assets: set[str] = set()
    if not transactions.empty:
        for asset_id, group in transactions.groupby("asset_id"):
            products = set(group["product"].dropna().map(_clean_text))
            if asset_id in composition_assets or products.intersection(composition_products):
                assigned_tx_assets.add(asset_id)

    missing_targets = tx_assets - assigned_tx_assets
    if missing_targets and composition.empty:
        _message(messages, "Warning", "Cross-sheet", "No target allocations are available for traded assets.")
    elif missing_targets:
        _message(messages, "Warning", "Cross-sheet", "Some traded assets may be outside target allocation: " + ", ".join(sorted(missing_targets)))

    if not composition.empty:
        unmatched = [p for p in sorted(composition["product"].dropna().unique()) if p not in known_products and p not in tx_assets and p not in price_assets]
        if unmatched:
            _message(messages, "Warning", "Cross-sheet", "Target allocation instruments not found in transactions/prices: " + ", ".join(unmatched))

    combined_names = pd.concat(
        [
            transactions[["isin", "product"]] if not transactions.empty else pd.DataFrame(columns=["isin", "product"]),
            prices[["isin", "product"]] if not prices.empty else pd.DataFrame(columns=["isin", "product"]),
        ],
        ignore_index=True,
    )
    combined_names = combined_names[(combined_names["isin"].astype(str).str.len() > 0) & (combined_names["product"].astype(str).str.len() > 0)]
    for isin, group in combined_names.groupby("isin"):
        products = sorted(set(group["product"].map(_clean_text)))
        if len(products) > 1:
            _message(messages, "Warning", "Cross-sheet", f"ISIN {isin} has inconsistent product names: " + " | ".join(products))

    if not transactions.empty and not prices.empty:
        first_tx = transactions["date"].min()
        latest_price = prices["date"].max()
        if pd.notna(first_tx) and pd.notna(latest_price) and latest_price < first_tx:
            _message(messages, "Error", "Cross-sheet", "Historical prices end before the first transaction date.")


def build_upload_template_bytes(include_examples: bool = False) -> bytes:
    wb = Workbook()
    ws_comp = wb.active
    ws_comp.title = SHEET_PORTFOLIO_COMPOSITION
    ws_prices = wb.create_sheet(SHEET_HISTORICAL_PRICES)
    ws_tx = wb.create_sheet(SHEET_TRANSACTIONS)

    _build_composition_sheet(ws_comp, include_examples)
    _build_prices_sheet(ws_prices, include_examples)
    _build_transactions_sheet(ws_tx, include_examples)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def _build_transactions_sheet(ws, include_examples: bool) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.cell(1, 1, "All Transactions")
    for col, header in enumerate(TRANSACTION_EXCEL_HEADERS, start=1):
        ws.cell(2, col, header)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A2:J2"
    action_validation = DataValidation(type="list", formula1='"Buy,Sell,Transfer"', allow_blank=True)
    ws.add_data_validation(action_validation)
    action_validation.add("F3:F10000")

    if include_examples:
        sample_rows = [
            ["DEGIRO IRELAND", "iShares Core MSCI World UCITS (Acc)", "IE00B4L5Y983", pd.Timestamp("2019-03-01"), "EAM", "Buy", 42.0, 48.50, 2037.00, 0.0],
            ["DEGIRO IRELAND", "iShares Core MSCI EM IMI UCITS (Acc)", "IE00BKM4GZ66", pd.Timestamp("2019-03-08"), "EAM", "Buy", 25.0, 22.95, 573.75, 0.0],
            ["DKB", "iShares Core MSCI World Small Cap UCITS (Acc)", "IE00BF4RFH31", pd.Timestamp("2020-01-10"), "XETRA", "Buy", 140.0, 4.15, 581.00, 1.50],
        ]
        for row_num, values in enumerate(sample_rows, start=3):
            for col, value in enumerate(values, start=1):
                ws.cell(row_num, col, value)

    _style_sheet(ws, max_col=10, percent_cols=[], date_cols=[4], money_cols=[8, 9, 10])
    widths = [18, 42, 16, 13, 12, 12, 12, 14, 17, 12]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws["A1"].comment = Comment("Only the All Transactions section is parsed. Cash-flow values should be entered as positive EUR amounts.", "Codex")


def _build_prices_sheet(ws, include_examples: bool) -> None:
    products = [
        ("iShares Core MSCI World UCITS (Acc)", "IE00B4L5Y983"),
        ("iShares Core MSCI EM IMI UCITS (Acc)", "IE00BKM4GZ66"),
        ("iShares Core MSCI World Small Cap UCITS (Acc)", "IE00BF4RFH31"),
    ]
    for block_idx, (product, isin) in enumerate(products):
        start = 1 + block_idx * 5
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + 3)
        ws.cell(1, start, product)
        for offset, header in enumerate(PRICE_HEADERS):
            ws.cell(2, start + offset, header)
        if include_examples:
            sample_prices = [
                (pd.Timestamp("2019-01-02"), [43.44, 21.92, 3.89][block_idx]),
                (pd.Timestamp("2019-03-01"), [48.50, 22.80, 4.02][block_idx]),
                (pd.Timestamp("2019-03-08"), [47.92, 22.95, 4.01][block_idx]),
                (pd.Timestamp("2020-01-10"), [58.20, 25.10, 4.15][block_idx]),
                (pd.Timestamp("2020-12-31"), [65.00, 28.50, 4.85][block_idx]),
            ]
            for row_num, (date, price) in enumerate(sample_prices, start=3):
                ws.cell(row_num, start, product)
                ws.cell(row_num, start + 1, isin)
                ws.cell(row_num, start + 2, date)
                ws.cell(row_num, start + 3, price)

    risk_free_start = 16
    ws.merge_cells(start_row=1, start_column=risk_free_start, end_row=1, end_column=risk_free_start + 7)
    ws.cell(1, risk_free_start, RISK_FREE_GROUP_HEADER)
    for offset, header in enumerate(RISK_FREE_EXCEL_HEADERS):
        ws.cell(2, risk_free_start + offset, header)
    if include_examples:
        risk_free_rows = [
            [2019, "Germany", -0.62, -0.60, -0.48, -0.19, 0.28, 0.37],
            [2020, "Germany", -0.73, -0.74, -0.72, -0.57, -0.18, -0.17],
            [2021, "Germany", -0.70, -0.62, -0.49, -0.18, 0.20, 0.34],
            [2022, "Germany", 2.83, 2.60, 2.30, 2.10, 2.18, 2.13],
            [2023, "Germany", 3.21, 2.57, 2.19, 2.02, 2.18, 2.10],
        ]
        for row_num, values in enumerate(risk_free_rows, start=3):
            for offset, value in enumerate(values):
                ws.cell(row_num, risk_free_start + offset, value / 100 if offset >= 2 else value)
    ws.freeze_panes = "A3"
    _style_sheet(
        ws,
        max_col=23,
        percent_cols=[18, 19, 20, 21, 22, 23],
        date_cols=[3, 8, 13],
        money_cols=[4, 9, 14],
    )
    for col in range(1, 16):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col % 5 else 4
    ws.column_dimensions["P"].width = 10
    ws.column_dimensions["Q"].width = 16
    for col in range(18, 24):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws["A1"].comment = Comment("Add more assets by continuing the four-column Product/ISIN/Date/Price block plus one blank spacer column. The optional government-bond block supplies risk-free rates for risk-adjusted metrics.", "Codex")


def _build_composition_sheet(ws, include_examples: bool) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws.cell(1, 1, SHEET_PORTFOLIO_COMPOSITION)
    headers = [
        "Year",
        "Instrument",
        COMPOSITION_TARGET_HEADER,
        "Instrument 2",
        COMPOSITION_TARGET_HEADER,
        "Instrument 3",
        COMPOSITION_TARGET_HEADER,
        COMPOSITION_TOLERANCE_HEADER,
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(2, col, header)
    if include_examples:
        rows = [
            [
                2019,
                "iShares Core MSCI World UCITS (Acc)",
                0.70,
                "iShares Core MSCI EM IMI UCITS (Acc)",
                0.20,
                "iShares Core MSCI World Small Cap UCITS (Acc)",
                0.10,
                0.05,
            ],
            [
                2020,
                "iShares Core MSCI World UCITS (Acc)",
                0.68,
                "iShares Core MSCI EM IMI UCITS (Acc)",
                0.22,
                "iShares Core MSCI World Small Cap UCITS (Acc)",
                0.10,
                0.05,
            ],
        ]
        for row_num, values in enumerate(rows, start=3):
            for col, value in enumerate(values, start=1):
                ws.cell(row_num, col, value)
    ws.freeze_panes = "A3"
    _style_sheet(ws, max_col=8, percent_cols=[3, 5, 7, 8], date_cols=[], money_cols=[])
    widths = [10, 42, 24, 42, 24, 42, 24, 26]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws["A1"].comment = Comment("Add more Instrument / Target Weight pairs before the final tolerance column.", "Codex")


def _style_sheet(ws, max_col: int, percent_cols: list[int], date_cols: list[int], money_cols: list[int]) -> None:
    title_fill = PatternFill("solid", fgColor="0F172A")
    header_fill = PatternFill("solid", fgColor="DBEAFE")
    title_font = Font(color="FFFFFF", bold=True, size=13)
    header_font = Font(color="111827", bold=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(bottom=thin)

    for col in range(1, max_col + 1):
        cell = ws.cell(1, col)
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header = ws.cell(2, col)
        header.fill = header_fill
        header.font = header_font
        header.border = border
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in percent_cols:
        for row in range(3, 500):
            ws.cell(row, col).number_format = "0.00%"
    for col in date_cols:
        for row in range(3, 500):
            ws.cell(row, col).number_format = "yyyy-mm-dd"
    for col in money_cols:
        for row in range(3, 500):
            ws.cell(row, col).number_format = '#,##0.00'
