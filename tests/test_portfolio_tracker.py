from __future__ import annotations

from io import BytesIO
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from portfolio_analytics.engine import (
    build_analysis,
    build_rolling_performance,
    risk_free_rate_for_period,
)
from portfolio_analytics.charts import rolling_metric_chart
from portfolio_analytics.models import PortfolioWorkbook
from portfolio_analytics.workbook import (
    REQUIRED_SHEETS,
    TRANSACTION_EXCEL_HEADERS,
    build_upload_template_bytes,
    parse_workbook,
)


def test_template_contains_exact_required_sheets_and_headers():
    data = build_upload_template_bytes(include_examples=False)
    workbook = load_workbook(BytesIO(data))

    assert workbook.sheetnames == REQUIRED_SHEETS
    composition = workbook["Portfolio Composition"]
    assert composition["A1"].value == "Portfolio Composition"
    assert "A1:K1" in {str(rng) for rng in composition.merged_cells.ranges}
    assert [composition.cell(2, col).value for col in range(1, 12)] == [
        "Year",
        "Instrument",
        "ISIN",
        "Target Weight in Portfolio",
        "Instrument 2",
        "ISIN",
        "Target Weight in Portfolio",
        "Instrument 3",
        "ISIN",
        "Target Weight in Portfolio",
        "Weight Deviation Tolerance",
    ]
    tx = workbook["Transactions"]
    assert tx["A1"].value == "All Transactions"
    assert "A1:J1" in {str(rng) for rng in tx.merged_cells.ranges}
    assert [tx.cell(2, col).value for col in range(1, 11)] == TRANSACTION_EXCEL_HEADERS


def test_prefilled_template_contains_2019_2026_isin_aware_composition_rows():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    composition = workbook["Portfolio Composition"]

    assert [composition.cell(row, 1).value for row in range(3, 11)] == list(range(2019, 2027))
    assert composition["C3"].value == "IE00B4L5Y983"
    assert composition["F3"].value == "IE00BKM4GZ66"
    assert composition["I3"].value == "IE00BF4RFH31"
    assert composition["D3"].value == 0.60
    assert composition["G3"].value == 0.40
    assert composition["J3"].value == 0.00
    assert composition["G7"].value == 0.30
    assert composition["J7"].value == 0.10


def test_sample_template_parses_and_builds_dynamic_analysis():
    parsed = parse_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    analysis = build_analysis(parsed)

    assert not any(message.severity == "Error" for message in parsed.validation)
    assert set(parsed.workbook_sheets) == set(REQUIRED_SHEETS)
    assert parsed.transactions["asset_id"].nunique() == 3
    assert parsed.prices["asset_id"].nunique() == 3
    assert list(parsed.composition.columns) == [
        "year",
        "product",
        "isin",
        "asset_id",
        "target_weight",
        "tolerance",
        "source_row",
        "target_column",
    ]
    assert len(parsed.composition) == 24
    first_year = parsed.composition[parsed.composition["year"].eq(2019)]
    assert set(first_year["asset_id"]) == {"IE00B4L5Y983", "IE00BKM4GZ66", "IE00BF4RFH31"}
    assert first_year.set_index("isin").loc["IE00B4L5Y983", "target_weight"] == 0.60
    assert first_year.set_index("isin").loc["IE00BKM4GZ66", "target_weight"] == 0.40
    assert first_year.set_index("isin").loc["IE00BF4RFH31", "target_weight"] == 0.00
    assert not analysis.daily_portfolio.empty
    assert analysis.daily_portfolio.iloc[0]["is_baseline"]
    assert analysis.daily_portfolio.iloc[0]["value_eur"] == 0
    assert analysis.allocation["target_weight"].sum() > 0


def test_risk_free_block_is_parsed_and_used_for_annual_metrics():
    parsed = parse_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    analysis = build_analysis(parsed)

    assert not parsed.risk_free_rates.empty
    assert list(parsed.risk_free_rates.columns) == [
        "year",
        "country",
        "yield_1y",
        "yield_3y",
        "yield_5y",
        "yield_10y",
        "yield_20y",
        "yield_30y",
    ]
    metric_2019 = analysis.annual.loc[analysis.annual["year"].eq(2019)].iloc[0]
    assert metric_2019["risk_free_rate"] == -0.0062
    assert "No risk-free-rate block found" not in " ".join(
        message.message for message in parsed.validation
    )


def test_risk_free_tie_breaks_prefer_the_next_longer_maturity():
    rates = pd.DataFrame(
        [{"year": 2024, "yield_1y": 0.01, "yield_3y": 0.03, "yield_5y": 0.05}]
    )

    assert risk_free_rate_for_period(rates, 2, 2024) == 0.03
    assert risk_free_rate_for_period(rates, 4, 2024) == 0.05


def test_price_block_identity_uses_most_common_isin_when_row_isins_vary():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    prices = workbook["Historical Prices"]
    for row_num, isin in zip(
        range(3, 8),
        ["IE00BF4RFH31", "IE00BF4RFH32", "IE00BF4RFH33", "IE00BF4RFH34", "IE00BF4RFH35"],
    ):
        prices.cell(row_num, 12, isin)

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))
    analysis = build_analysis(parsed)

    small_cap_tx = parsed.transactions[parsed.transactions["product"].str.contains("Small Cap")].iloc[0]
    small_cap_prices = parsed.prices[parsed.prices["asset_id"].eq(small_cap_tx["asset_id"])]
    small_cap_holding = analysis.holdings[analysis.holdings["product"].str.contains("Small Cap")].iloc[0]

    assert small_cap_tx["asset_id"] == "IE00BF4RFH31"
    assert small_cap_prices["isin"].nunique() == 1
    assert set(small_cap_prices["isin"]) == {"IE00BF4RFH31"}
    assert small_cap_prices["asset_id"].nunique() == 1
    assert small_cap_holding["quantity"] == 140
    assert small_cap_holding["price"] == 4.85
    assert abs(small_cap_holding["value_eur"] - 679.0) < 0.001
    assert any("contains inconsistent ISIN values" in message.message for message in parsed.validation)


def test_text_transaction_dates_are_rejected_by_default():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    transactions = workbook["Transactions"]
    transactions.cell(3, 4, "25/6/2026")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))

    assert 3 not in set(parsed.transactions["source_row"])
    assert any(
        "Transactions row 3 has a text date '25/6/2026'" in message.message
        and message.severity == "Error"
        for message in parsed.validation
    )


def test_permissive_text_transaction_dates_parse_day_first_with_warning():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    transactions = workbook["Transactions"]
    transactions.cell(3, 4, "25/6/2026")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)), text_date_mode="permissive")
    parsed_row = parsed.transactions[parsed.transactions["source_row"].eq(3)].iloc[0]

    assert parsed_row["date"] == pd.Timestamp("2026-06-25")
    assert any("Parsed day-first in permissive mode" in message.message for message in parsed.validation)


def test_transaction_product_isin_mismatch_warns_precisely():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    transactions = workbook["Transactions"]
    transactions.cell(3, 3, "IE00BKM4GZ66")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))

    assert any(
        "Transaction row 3 has a product/ISIN mismatch" in message.message
        and "ISIN matching was used" in message.message
        and message.rows == "3"
        for message in parsed.validation
    )


def test_product_name_can_differ_when_isin_matches():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    transactions = workbook["Transactions"]
    transactions.cell(3, 2, "World ETF alias")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))
    analysis = build_analysis(parsed)

    assert parsed.transactions.loc[parsed.transactions["source_row"].eq(3), "asset_id"].iloc[0] == "IE00B4L5Y983"
    assert not analysis.daily_portfolio.empty
    assert any(
        "Product names differ for ISIN IE00B4L5Y983; using ISIN as the asset key" in message.message
        for message in parsed.validation
    )


def test_same_product_with_different_isins_warns_across_sheets():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    composition = workbook["Portfolio Composition"]
    for row_num in range(3, 11):
        composition.cell(row_num, 3, "IE00DIFFERENT1")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))

    assert any(
        "has inconsistent ISINs across sheets" in message.message
        and "IE00B4L5Y983" in message.message
        and "IE00DIFFERENT1" in message.message
        for message in parsed.validation
    )


def test_allocation_uses_isin_based_targets_when_composition_product_differs():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    composition = workbook["Portfolio Composition"]
    for row_num in range(3, 11):
        composition.cell(row_num, 2, "World target alias")

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))
    analysis = build_analysis(parsed)

    world_target = analysis.allocation[analysis.allocation["asset_id"].eq("IE00B4L5Y983")].iloc[0]
    assert world_target["target_weight"] == 0.60
    assert world_target["tolerance"] == 0.05


def test_blank_composition_isin_warns_and_falls_back_to_product_matching():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    composition = workbook["Portfolio Composition"]
    for row_num in range(3, 11):
        composition.cell(row_num, 3).value = None

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))
    analysis = build_analysis(parsed)

    world_target = analysis.allocation[analysis.allocation["asset_id"].eq("IE00B4L5Y983")].iloc[0]
    assert world_target["target_weight"] == 0.60
    assert any(
        "ISIN is blank for a target allocation row" in message.message
        for message in parsed.validation
    )


def test_suspicious_transaction_price_warning_includes_reference_price_details():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    transactions = workbook["Transactions"]
    transactions.cell(3, 8, 9.16)

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))

    assert any(
        "Transaction row 3 (2019-03-01, Buy" in message.message
        and "transaction price 9.16" in message.message
        and "nearest prior historical price 48.50 on 2019-03-01" in message.message
        and "Threshold is 20.0%" in message.message
        for message in parsed.validation
    )


def test_daily_performance_reconciliation_reports_mismatches():
    workbook = load_workbook(BytesIO(build_upload_template_bytes(include_examples=True)))
    daily = workbook.create_sheet("Daily Performance")
    headers = [
        "Date",
        "Portfolio Value in €",
        "Total Investment in €",
        "Total Fees in €",
        "Cash Flow in €",
        "Return in €",
        "Total Return in €",
        "HPR",
        "1+HPR",
        "TWR",
    ]
    for col, header in enumerate(headers, start=1):
        daily.cell(1, col, header)
    for col, value in enumerate([pd.Timestamp("2020-12-31"), 0, 0, 0, 0, 0, 0, 0, 1, 0], start=1):
        daily.cell(2, col, value)

    parsed = parse_workbook(BytesIO(_workbook_bytes(workbook)))
    analysis = build_analysis(parsed)

    assert not analysis.reconciliation.empty
    assert "Mismatch" in set(analysis.reconciliation["Status"])
    assert any(
        "App output does not match workbook Daily Performance" in message.message
        for message in parsed.validation
    )


def test_transfer_is_preserved_but_ignored_for_quantity():
    workbook = _workbook_from_frames(
        transactions=[
            ["Broker", "Asset A", "IE0000000001", "2020-01-01", "XETRA", "Buy", 10, 10, 100, 1, 3],
            ["Broker", "Asset A", "IE0000000001", "2020-01-02", "XETRA", "Transfer", 100, 10, 1000, 0, 4],
        ],
        prices=[
            ["Asset A", "IE0000000001", "2020-01-01", 10, 3, 1],
            ["Asset A", "IE0000000001", "2020-01-02", 11, 4, 1],
        ],
    )
    analysis = build_analysis(workbook)

    latest = analysis.holdings.iloc[0]
    assert latest["quantity"] == 10
    assert latest["value_eur"] == 110


def test_sell_decreases_quantity_and_invested_capital():
    workbook = _workbook_from_frames(
        transactions=[
            ["Broker", "Asset A", "IE0000000001", "2020-01-01", "XETRA", "Buy", 10, 10, 100, 0, 3],
            ["Broker", "Asset A", "IE0000000001", "2020-01-02", "XETRA", "Sell", 4, 12, 48, 1, 4],
        ],
        prices=[
            ["Asset A", "IE0000000001", "2020-01-01", 10, 3, 1],
            ["Asset A", "IE0000000001", "2020-01-02", 12, 4, 1],
        ],
    )
    analysis = build_analysis(workbook)

    latest = analysis.holdings.iloc[0]
    assert latest["quantity"] == 6
    assert latest["total_investment_eur"] == 53
    assert latest["total_fees_eur"] == 1


def test_rolling_5y_uses_inclusive_end_year():
    dates = pd.to_datetime([f"{year}-12-31" for year in range(2019, 2024)])
    daily = pd.DataFrame(
        {
            "date": dates,
            "value_eur": [100, 110, 125, 130, 150],
            "total_investment_eur": [100, 100, 100, 100, 100],
            "total_fees_eur": [0, 0, 0, 0, 0],
            "cash_flow_eur": [0, 0, 0, 0, 0],
            "return_eur": [0, 10, 15, 5, 20],
            "total_return_eur": [0, 10, 25, 30, 50],
            "hpr": [0, 0.10, 0.13636, 0.04, 0.15384],
            "one_plus_hpr": [1, 1.10, 1.13636, 1.04, 1.15384],
            "twr": [0, 0.10, 0.25, 0.30, 0.50],
            "drawdown": [0, 0, 0, 0, 0],
            "drawdown_pct": [0, 0, 0, 0, 0],
            "is_baseline": [False, False, False, False, False],
        }
    )

    rolling = build_rolling_performance(daily, annual_risk_free_rate=0.0)

    assert "2019-2023" in set(rolling["period"])
    row = rolling[rolling["period"] == "2019-2023"].iloc[0]
    assert row["start_year"] == 2019
    assert row["end_year"] == 2023


def test_rolling_metric_chart_uses_plotly_compatible_discrete_colors():
    rolling = pd.DataFrame(
        {
            "period": ["2019-2019", "2020-2020", "2019-2020"],
            "period_years": [1, 1, 2],
            "annualized_twr": [0.05, 0.08, 0.065],
        }
    )

    fig = rolling_metric_chart(rolling, "annualized_twr", "Rolling annualized TWR")

    assert len(fig.data) == 2
    assert {trace.name for trace in fig.data} == {"1Y", "2Y"}


def _workbook_bytes(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _workbook_from_frames(transactions: list[list], prices: list[list]) -> PortfolioWorkbook:
    tx = pd.DataFrame(
        transactions,
        columns=[
            "broker",
            "product",
            "isin",
            "date",
            "exchange",
            "action",
            "quantity",
            "price_eur",
            "cash_flow_eur",
            "fee_eur",
            "source_row",
        ],
    )
    tx["date"] = pd.to_datetime(tx["date"])
    tx["asset_id"] = tx["isin"]
    px = pd.DataFrame(prices, columns=["product", "isin", "date", "price", "source_row", "block"])
    px["date"] = pd.to_datetime(px["date"])
    px["asset_id"] = px["isin"]
    composition = pd.DataFrame(
        [
            {
                "year": 2020,
                "product": "Asset A",
                "isin": "IE0000000001",
                "asset_id": "IE0000000001",
                "target_weight": 1.0,
                "tolerance": 0.05,
                "source_row": 3,
                "target_column": "D",
            }
        ]
    )
    return PortfolioWorkbook(
        transactions=tx,
        prices=px,
        composition=composition,
        validation=[],
        workbook_sheets=REQUIRED_SHEETS,
    )
